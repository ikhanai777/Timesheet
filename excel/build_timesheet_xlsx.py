"""Builds the Timesheet Excel workbook (same rules as the Android app, see docs/SPEC.md).

Usage: python3 excel/build_timesheet_xlsx.py [year] [output.xlsx]

Sheets: Guide, Settings, Projects, <year> Dashboard, and one sheet per month. Each month sheet has
an entry log (left) and a dashboard (right): month KPIs, weekly summary, chart, projects this
month and a daily summary. Every total is a formula, so the workbook recalculates as you type.
"""
import calendar
import datetime as dt
import sys

from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.comments import Comment
from openpyxl.formatting.rule import FormulaRule
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.worksheet.datavalidation import DataValidation

YEAR = int(sys.argv[1]) if len(sys.argv) > 1 else 2026
OUT = sys.argv[2] if len(sys.argv) > 2 else f"excel/Timesheet_{YEAR}.xlsx"

LOG_FIRST, LOG_ROWS = 7, 250
LOG_LAST = LOG_FIRST + LOG_ROWS - 1
PROJ_FIRST, PROJ_ROWS = 5, 200  # Projects sheet rows
PROJ_LAST = PROJ_FIRST + PROJ_ROWS - 1
HELP_FIRST = LOG_FIRST  # month-sheet helper rows, aligned 1:1 with Projects rows
HELP_LAST = HELP_FIRST + PROJ_ROWS - 1
WEEK_FIRST = 10  # weekly rows 10..15
TOP_FIRST, TOP_N = 37, 15
DAILY_HEADER = 57
DAILY_FIRST = 58
MONTHS = [calendar.month_name[m] for m in range(1, 13)]

# ---- Styles -------------------------------------------------------------------------------
FONT = "Arial"
NAVY, BLUE, ORANGE, GREY = "1F4E8C", "1F6FEB", "D9730D", "8C959F"
f_base = Font(name=FONT, size=10)
f_bold = Font(name=FONT, size=10, bold=True)
f_title = Font(name=FONT, size=18, bold=True, color=NAVY)
f_sub = Font(name=FONT, size=10, italic=True, color="555555")
f_head = Font(name=FONT, size=10, bold=True, color="FFFFFF")
f_section = Font(name=FONT, size=11, bold=True, color=NAVY)
f_input = Font(name=FONT, size=10, color="0000FF")
f_kpi_label = Font(name=FONT, size=9, bold=True, color="555555")
f_kpi = Font(name=FONT, size=16, bold=True, color="1A1A1A")
fill_head = PatternFill("solid", fgColor=NAVY)
fill_input = PatternFill("solid", fgColor="FFF9DB")
fill_kpi = PatternFill("solid", fgColor="EEF2F8")
fill_total = PatternFill("solid", fgColor="DCE3EE")
fill_ntime = PatternFill("solid", fgColor=BLUE)
fill_p6 = PatternFill("solid", fgColor=ORANGE)
fill_oh = PatternFill("solid", fgColor=GREY)
thin = Side(style="thin", color="C8CED8")
box = Border(left=thin, right=thin, top=thin, bottom=thin)
HOURS = '0.0;(0.0);"-"'
PCT = '0%;(0%);"-"'
center = Alignment(horizontal="center", vertical="center")
wrap = Alignment(wrap_text=True, vertical="top")


def style_range(ws, ref, font=None, fill=None, fmt=None, border=True, align=None):
    for row in ws[ref]:
        for c in row:
            if font:
                c.font = font
            if fill:
                c.fill = fill
            if fmt:
                c.number_format = fmt
            if border:
                c.border = box
            if align:
                c.alignment = align


def header_row(ws, row, col, labels, fills=None):
    for i, label in enumerate(labels):
        c = ws.cell(row=row, column=col + i, value=label)
        c.font = f_head
        c.fill = (fills or {}).get(label, fill_head)
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        c.border = box


def section(ws, ref, text):
    ws[ref] = text
    ws[ref].font = f_section


def widths(ws, mapping):
    for col, w in mapping.items():
        ws.column_dimensions[col].width = w


wb = Workbook()
wb.remove(wb.active)

# ---- Guide --------------------------------------------------------------------------------
guide = wb.create_sheet("Guide")
guide.sheet_view.showGridLines = False
widths(guide, {"A": 2, "B": 14, "C": 13, "D": 24, "E": 14, "F": 14, "G": 34, "H": 9, "I": 8})
guide["B1"] = f"Timesheet {YEAR}"
guide["B1"].font = f_title
guide["B2"] = "Log daily project hours against NTIME or P6. Unbooked hours become Overhead (Department Hours) automatically."
guide["B2"].font = f_sub
lines = [
    ("How to use", None),
    ("1.", "Settings: check available hours per day (default 8), week start day, where overhead is booked, and overtime."),
    ("2.", "Projects: enter each project once (name, project number, Oracle code). They then appear in the dropdowns."),
    ("3.", "Month sheets: add one row per task in the Entry log - date, project (dropdown), task description, NTIME/P6, hours."),
    ("", "Project number and Oracle code fill in automatically. The Check column flags anything missing or a day over its hours."),
    ("4.", "Daily summary (bottom right of each month): change Day type for leave, public holidays, or a day of only overhead."),
    ("5.", "Everything else is calculated: month KPIs, weekly summary and chart, projects this month, and the year dashboard."),
    ("", ""),
    ("Rules", None),
    ("Overhead", "On a Working day with project hours: overhead = available hours - project hours (never below 0)."),
    ("", "Example: Project A 4 h in NTIME + Project B 3 h in P6 = 7 h, so 1 h goes to Overhead."),
    ("Overhead only", "Set Day type to 'Overhead only' for a working day with no project work: the whole day is overhead."),
    ("Not logged", "A Working day up to today with no entries is shown as 'Not logged' and counts 0 hours."),
    ("Leave etc.", "Leave, Public Holiday and Weekend days count no overhead."),
    ("Overtime", "If Allow overtime is No, a day over its available hours is flagged in red. Extra hours show as Overtime."),
    ("Weeks", "Weeks start on the day chosen in Settings and are cut at the month's first and last day."),
    ("", ""),
    ("Legend", None),
]
r = 4
for a, b in lines:
    if b is None:
        guide.cell(row=r, column=2, value=a).font = f_section
    else:
        guide.cell(row=r, column=2, value=a).font = f_bold
        guide.cell(row=r, column=3, value=b).font = f_base
    r += 1
guide.cell(row=r, column=2).fill = fill_input
guide.cell(row=r, column=2).border = box
guide.cell(row=r, column=3, value="Yellow cells with blue text are inputs - type or pick from the dropdown.").font = f_base
r += 1
guide.cell(row=r, column=2, value="1.0").font = f_base
guide.cell(row=r, column=2).border = box
guide.cell(row=r, column=3, value="White cells are formulas - don't type over them.").font = f_base
r += 2
guide.cell(row=r, column=2, value="Example entry row (format to follow)").font = f_section
r += 1
header_row(guide, r, 2, ["Date", "Day", "Project Name", "Project Number", "Oracle Code", "Task Description", "System", "Hours"])
r += 1
example = [dt.date(YEAR, 1, 5), "Mon", "Project A", "PRJ-2041", "4100.22.07", "Reviewed piping layout drawings", "NTIME", 4]
for i, v in enumerate(example):
    c = guide.cell(row=r, column=2 + i, value=v)
    c.font = f_input if i in (0, 2, 5, 6, 7) else f_base
    c.fill = fill_input if i in (0, 2, 5, 6, 7) else PatternFill()
    c.border = box
guide.cell(row=r, column=2).number_format = "dd-mmm-yy"
guide.cell(row=r, column=9).number_format = HOURS

# ---- Settings -----------------------------------------------------------------------------
st = wb.create_sheet("Settings")
st.sheet_view.showGridLines = False
widths(st, {"A": 2, "B": 30, "C": 14, "D": 70})
st["B1"] = "Settings"
st["B1"].font = f_title
header_row(st, 4, 2, ["Setting", "Value", "Notes"])
settings = [
    ("Year", YEAR, "Fixed for this workbook. Rebuild with excel/build_timesheet_xlsx.py for another year.", False),
    ("Available hours per day", 8, "Hours in a working day. Overhead = available hours - project hours. Can be overridden per day in each month's daily summary.", True),
    ("Week starts on", "Monday", "Monday, Sunday or Saturday. Used by the weekly summaries.", True),
    ("Overhead booked in", "None", "None keeps overhead separate. NTIME or P6 adds overhead to that system's totals.", True),
    ("Allow overtime", "No", "No: a day over its available hours is flagged in red. Yes: extra hours are shown as Overtime.", True),
    ("Week start (number)", '=MATCH(C7,{"Monday","Tuesday","Wednesday","Thursday","Friday","Saturday","Sunday"},0)',
     "Calculated - used by the formulas.", False),
]
for i, (label, value, note, editable) in enumerate(settings):
    row = 5 + i
    st.cell(row=row, column=2, value=label).font = f_bold
    c = st.cell(row=row, column=3, value=value)
    c.font = f_input if editable else f_base
    c.fill = fill_input if editable else PatternFill()
    c.alignment = center
    st.cell(row=row, column=4, value=note).font = f_base
    for col in (2, 3, 4):
        st.cell(row=row, column=col).border = box
dv = DataValidation(type="decimal", operator="between", formula1="1", formula2="24", showErrorMessage=True,
                    error="Enter hours between 1 and 24.")
st.add_data_validation(dv)
dv.add("C6")
for ref, options in (("C7", "Monday,Sunday,Saturday"), ("C8", "None,NTIME,P6"), ("C9", "No,Yes")):
    v = DataValidation(type="list", formula1=f'"{options}"', allow_blank=False)
    st.add_data_validation(v)
    v.add(ref)
for name, ref in (("AvailHours", "$C$6"), ("WeekStartNum", "$C$10"), ("OHSystem", "$C$8"), ("AllowOT", "$C$9")):
    wb.defined_names[name] = DefinedName(name, attr_text=f"Settings!{ref}")

# ---- Projects -----------------------------------------------------------------------------
pj = wb.create_sheet("Projects")
pj.sheet_view.showGridLines = False
widths(pj, {"A": 2, "B": 34, "C": 16, "D": 16, "E": 12, "F": 12, "G": 12, "H": 24, "I": 4})
pj["B1"] = "Projects"
pj["B1"].font = f_title
pj["B2"] = "Enter each project once. Names must be unique - they are what you pick in the month sheets."
pj["B2"].font = f_sub
header_row(pj, 4, 2, ["Project Name", "Project Number", "Oracle Code", f"NTIME {YEAR}", f"P6 {YEAR}", f"Total {YEAR}", "Check"],
           {f"NTIME {YEAR}": fill_ntime, f"P6 {YEAR}": fill_p6})
for row in range(PROJ_FIRST, PROJ_LAST + 1):
    h = row - PROJ_FIRST + HELP_FIRST
    for col in "BCD":
        c = pj[f"{col}{row}"]
        c.fill, c.font, c.border = fill_input, f_input, box
    pj[f"E{row}"] = f'=IF(B{row}="","",' + "+".join(f"{m}!X{h}" for m in MONTHS) + ")"
    pj[f"F{row}"] = f'=IF(B{row}="","",' + "+".join(f"{m}!Y{h}" for m in MONTHS) + ")"
    pj[f"G{row}"] = f'=IF(B{row}="","",E{row}+F{row})'
    pj[f"H{row}"] = (f'=IF(B{row}="","",IF(COUNTIF(ProjNames,B{row})>1,"Duplicate name!",'
                     f'IF(OR(C{row}="",D{row}=""),"Add number and code","✓")))')
    pj[f"I{row}"] = f"=IF(N(G{row})>0,G{row}+({PROJ_LAST + 1}-ROW())/10^6,0)"
    for col in "EFG":
        pj[f"{col}{row}"].number_format = HOURS
    for col in "EFGH":
        pj[f"{col}{row}"].font, pj[f"{col}{row}"].border = f_base, box
pj["B5"], pj["C5"], pj["D5"] = "Project A", "PRJ-2041", "4100.22.07"
pj["B5"].comment = Comment("Example project - overwrite it with your own.", "Timesheet")
pj.column_dimensions["I"].hidden = True
pj.freeze_panes = "B5"
for name, col in (("ProjNames", "B"), ("ProjNumbers", "C"), ("ProjCodes", "D")):
    wb.defined_names[name] = DefinedName(name, attr_text=f"Projects!${col}${PROJ_FIRST}:${col}${PROJ_LAST}")
pj.conditional_formatting.add(f"H{PROJ_FIRST}:H{PROJ_LAST}", FormulaRule(
    formula=[f'AND(H{PROJ_FIRST}<>"",H{PROJ_FIRST}<>"✓")'], font=Font(name=FONT, color="C00000", bold=True)))
pj.conditional_formatting.add(f"B{PROJ_FIRST}:B{PROJ_LAST}", FormulaRule(
    formula=[f'COUNTIF(ProjNames,B{PROJ_FIRST})>1'], fill=PatternFill("solid", fgColor="F8D7DA")))

# ---- Year dashboard (filled after the month sheets exist) ---------------------------------
yd = wb.create_sheet(f"{YEAR} Dashboard")


# ---- Month sheets -------------------------------------------------------------------------
def build_month(m):
    name = MONTHS[m - 1]
    ws = wb.create_sheet(name)
    ws.sheet_view.showGridLines = False
    ws.sheet_properties.tabColor = NAVY
    days = calendar.monthrange(YEAR, m)[1]
    first, last = dt.date(YEAR, m, 1), dt.date(YEAR, m, days)
    d_last = DAILY_FIRST + days - 1
    LOG = lambda col: f"${col}${LOG_FIRST}:${col}${LOG_LAST}"  # noqa: E731
    DAY = lambda col: f"${col}${DAILY_FIRST}:${col}${d_last}"  # noqa: E731

    widths(ws, {"A": 2, "B": 11, "C": 6, "D": 26, "E": 14, "F": 13, "G": 38, "H": 9, "I": 8, "J": 24, "K": 3,
                "L": 24, "M": 13, "N": 15, "O": 12, "P": 12, "Q": 12, "R": 12, "S": 12, "T": 10, "U": 24})
    ws["B1"] = f"{name} {YEAR}"
    ws["B1"].font = f_title
    ws["B2"] = "Add one row per task in the Entry log. Everything on the right updates automatically."
    ws["B2"].font = f_sub

    # Hidden helper cells: month start, month end, first week start.
    ws["AC1"], ws["AC2"], ws["AC3"] = "Helpers", first, last
    ws["AC4"] = "=AC2-MOD(WEEKDAY(AC2,2)-WeekStartNum,7)"
    for ref in ("AC2", "AC3", "AC4"):
        ws[ref].number_format = "dd-mmm-yy"

    # Entry log ------------------------------------------------------------------------
    section(ws, "B3", "ENTRY LOG")
    header_row(ws, 6, 2, ["Date", "Day", "Project Name", "Project Number", "Oracle Code", "Task Description",
                          "System", "Hours", "Check"])
    for row in range(LOG_FIRST, LOG_LAST + 1):
        for col in "BDGHI":
            c = ws[f"{col}{row}"]
            c.fill, c.font, c.border = fill_input, f_input, box
        ws[f"B{row}"].number_format = "dd-mmm-yy"
        ws[f"I{row}"].number_format = "0.0#"
        ws[f"G{row}"].alignment = Alignment(wrap_text=False)
        ws[f"C{row}"] = f'=IF(B{row}="","",TEXT(B{row},"ddd"))'
        ws[f"E{row}"] = f'=IF(D{row}="","",IFERROR(INDEX(ProjNumbers,MATCH(D{row},ProjNames,0)),""))'
        ws[f"F{row}"] = f'=IF(D{row}="","",IFERROR(INDEX(ProjCodes,MATCH(D{row},ProjNames,0)),""))'
        avail = f"INDEX({DAY('O')},DAY(B{row}))"
        ws[f"J{row}"] = (
            f'=IF(COUNTA(B{row},D{row},G{row},H{row},I{row})=0,"",'
            f'IF(B{row}="","Enter date",'
            f'IF(OR(B{row}<$AC$2,B{row}>$AC$3),"Date not in {name}",'
            f'IF(D{row}="","Pick project",'
            f'IF(ISNA(MATCH(D{row},ProjNames,0)),"Project not in list",'
            f'IF(G{row}="","Add task description",'
            f'IF(H{row}="","Pick NTIME or P6",'
            f'IF(N(I{row})<=0,"Enter hours",'
            f'IF(AND(AllowOT<>"Yes",SUMIF({LOG("B")},B{row},{LOG("I")})>{avail}),'
            f'"Day over "&{avail}&" h","✓")))))))))'
        )
        for col in "CEFJ":
            ws[f"{col}{row}"].font, ws[f"{col}{row}"].border = f_base, box
        ws[f"C{row}"].alignment = center
    log_ref = f"B{LOG_FIRST}:J{LOG_LAST}"
    dv_date = DataValidation(type="date", operator="between", formula1=f"DATE({YEAR},{m},1)",
                             formula2=f"DATE({YEAR},{m},{days})", showErrorMessage=True,
                             error=f"Enter a date in {name} {YEAR}.", errorTitle="Date")
    dv_proj = DataValidation(type="list", formula1="=ProjNames", showErrorMessage=True,
                             error="Pick a project from the list. Add new projects on the Projects sheet first.",
                             errorTitle="Project")
    dv_sys = DataValidation(type="list", formula1='"NTIME,P6"', showErrorMessage=True)
    dv_hours = DataValidation(type="decimal", operator="between", formula1="0.01", formula2="24",
                              showErrorMessage=True, error="Hours must be more than 0 and at most 24.")
    dv_desc = DataValidation(type="textLength", operator="lessThanOrEqual", formula1="200",
                             showErrorMessage=True, error="Keep the task description to 200 characters.")
    for v, col in ((dv_date, "B"), (dv_proj, "D"), (dv_sys, "H"), (dv_hours, "I"), (dv_desc, "G")):
        ws.add_data_validation(v)
        v.add(f"{col}{LOG_FIRST}:{col}{LOG_LAST}")
    ws.conditional_formatting.add(f"J{LOG_FIRST}:J{LOG_LAST}", FormulaRule(
        formula=[f'AND(J{LOG_FIRST}<>"",J{LOG_FIRST}<>"✓")'],
        font=Font(name=FONT, color="9C5700", bold=True), fill=PatternFill("solid", fgColor="FFEB9C")))
    ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LOG_LAST}", FormulaRule(
        formula=[f'H{LOG_FIRST}="NTIME"'], font=Font(name=FONT, color=BLUE, bold=True)))
    ws.conditional_formatting.add(f"H{LOG_FIRST}:H{LOG_LAST}", FormulaRule(
        formula=[f'H{LOG_FIRST}="P6"'], font=Font(name=FONT, color=ORANGE, bold=True)))

    # Month KPIs -------------------------------------------------------------------------
    section(ws, "L3", "MONTH SUMMARY")
    kpis = [
        ("Total hours", f"=SUM({DAY('T')})", HOURS),
        ('=IF(OHSystem="NTIME","NTIME (incl. overhead)","NTIME")',
         f'=SUM({DAY("P")})+IF(OHSystem="NTIME",SUM({DAY("R")}),0)', HOURS),
        ('=IF(OHSystem="P6","P6 (incl. overhead)","P6")', f'=SUM({DAY("Q")})+IF(OHSystem="P6",SUM({DAY("R")}),0)', HOURS),
        ("Overhead", f"=SUM({DAY('R')})", HOURS),
        ("Overtime", f"=SUM({DAY('S')})", HOURS),
        ("Utilisation", f'=IF(L5=0,"-",(SUM({DAY("P")})+SUM({DAY("Q")}))/L5)', PCT),
        ("Days logged", f'=SUMPRODUCT(({DAY("N")}="Working")*(({DAY("P")}+{DAY("Q")})>0))+COUNTIF({DAY("N")},"Overhead only")', "0"),
        ("Days not logged", f'=COUNTIF({DAY("U")},"Not logged")', "0"),
    ]
    for i, (label, formula, fmt) in enumerate(kpis):
        col = get_column_letter(12 + i)
        lc, vc = ws[f"{col}4"], ws[f"{col}5"]
        lc.value, vc.value = label, formula
        lc.font, vc.font = f_kpi_label, f_kpi
        lc.fill = vc.fill = fill_kpi
        lc.alignment = Alignment(horizontal="center", wrap_text=True)
        vc.alignment = center
        vc.number_format = fmt
        lc.border = vc.border = box
    ws["M5"].font = Font(name=FONT, size=16, bold=True, color=BLUE)
    ws["N5"].font = Font(name=FONT, size=16, bold=True, color=ORANGE)
    ws["O5"].font = Font(name=FONT, size=16, bold=True, color="6E7781")
    ws.row_dimensions[4].height = 26
    ws.row_dimensions[5].height = 30
    ws.conditional_formatting.add("S5", FormulaRule(formula=["S5>0"], font=Font(name=FONT, size=16, bold=True, color="C00000")))

    # Weekly summary ---------------------------------------------------------------------
    section(ws, "L8", "WEEKLY SUMMARY")
    header_row(ws, 9, 12, ["Week", "From", "To", "NTIME", "P6", "Overhead", "Overtime", "Total"],
               {"NTIME": fill_ntime, "P6": fill_p6, "Overhead": fill_oh})
    ws["AF9"], ws["AG9"], ws["AH9"] = "NTIME", "P6", "Overhead"  # chart data: project-only system hours
    for k in range(1, 7):
        row = WEEK_FIRST + k - 1
        ws[f"M{row}"] = f'=IF($AC$4+7*{k - 1}>$AC$3,"",MAX($AC$2,$AC$4+7*{k - 1}))'
        ws[f"N{row}"] = f'=IF(M{row}="","",MIN($AC$3,$AC$4+7*{k}-1))'
        ws[f"L{row}"] = f'="Week {k}"&IF(M{row}="",""," ("&DAY(M{row})&"–"&DAY(N{row})&")")'
        rng = f'{DAY("L")},">="&$M{row},{DAY("L")},"<="&$N{row}'
        ws[f"AF{row}"] = f'=IF($M{row}="",0,SUMIFS({DAY("P")},{rng}))'
        ws[f"AG{row}"] = f'=IF($M{row}="",0,SUMIFS({DAY("Q")},{rng}))'
        ws[f"AH{row}"] = f'=IF($M{row}="",0,SUMIFS({DAY("R")},{rng}))'
        ws[f"O{row}"] = f'=IF($M{row}="","",AF{row}+IF(OHSystem="NTIME",AH{row},0))'
        ws[f"P{row}"] = f'=IF($M{row}="","",AG{row}+IF(OHSystem="P6",AH{row},0))'
        ws[f"Q{row}"] = f'=IF($M{row}="","",AH{row})'
        ws[f"R{row}"] = f'=IF($M{row}="","",SUMIFS({DAY("S")},{rng}))'
        ws[f"S{row}"] = f'=IF($M{row}="","",SUMIFS({DAY("T")},{rng}))'
    style_range(ws, f"L{WEEK_FIRST}:S{WEEK_FIRST + 5}", font=f_base, fmt=HOURS)
    style_range(ws, f"M{WEEK_FIRST}:N{WEEK_FIRST + 5}", fmt="ddd dd-mmm")
    tr = WEEK_FIRST + 6
    ws[f"L{tr}"] = "Month total"
    for col in "OPQRS":
        ws[f"{col}{tr}"] = f"=SUM({col}{WEEK_FIRST}:{col}{WEEK_FIRST + 5})"
    for col in ("AF", "AG", "AH"):
        ws[f"{col}{tr}"] = f"=SUM({col}{WEEK_FIRST}:{col}{WEEK_FIRST + 5})"
    style_range(ws, f"L{tr}:S{tr}", font=f_bold, fill=fill_total, fmt=HOURS)

    chart = BarChart()
    chart.type, chart.grouping, chart.overlap = "col", "stacked", 100
    chart.title = f"Hours per week - {name}"
    chart.y_axis.title = "Hours"
    chart.y_axis.majorGridlines = None
    chart.height, chart.width = 7.5, 22
    chart.add_data(Reference(ws, min_col=32, max_col=34, min_row=9, max_row=WEEK_FIRST + 5), titles_from_data=True)
    chart.set_categories(Reference(ws, min_col=12, min_row=WEEK_FIRST, max_row=WEEK_FIRST + 5))
    for series, color in zip(chart.series, (BLUE, ORANGE, GREY)):
        series.graphicalProperties.solidFill = color
        series.graphicalProperties.line.solidFill = color
    chart.legend.position = "b"
    chart.visible_cells_only = False  # its data lives in hidden helper columns
    ws.add_chart(chart, "L18")

    # Projects this month (top 15 by hours) ----------------------------------------------
    section(ws, "L35", "PROJECTS THIS MONTH")
    header_row(ws, 36, 12, ["Project", "Project Number", "Oracle Code", "NTIME", "P6", "Total", "% of month"],
               {"NTIME": fill_ntime, "P6": fill_p6})
    for i in range(PROJ_ROWS):
        h, p = HELP_FIRST + i, PROJ_FIRST + i
        ws[f"W{h}"] = f'=IF(Projects!B{p}="","",Projects!B{p})'
        ws[f"X{h}"] = f'=IF(W{h}="",0,SUMIFS({LOG("I")},{LOG("D")},W{h},{LOG("H")},"NTIME"))'
        ws[f"Y{h}"] = f'=IF(W{h}="",0,SUMIFS({LOG("I")},{LOG("D")},W{h},{LOG("H")},"P6"))'
        ws[f"Z{h}"] = f"=X{h}+Y{h}"
        ws[f"AA{h}"] = f"=IF(Z{h}>0,Z{h}+({PROJ_ROWS + 1}-{i + 1})/10^6,0)"
    keys = f"$AA${HELP_FIRST}:$AA${HELP_LAST}"
    for k in range(1, TOP_N + 1):
        row = TOP_FIRST + k - 1
        ws[f"AD{row}"] = f"=IF(LARGE({keys},{k})<=0,0,MATCH(LARGE({keys},{k}),{keys},0))"
        idx = f"$AD{row}"
        ws[f"L{row}"] = f'=IF({idx}=0,"",INDEX($W${HELP_FIRST}:$W${HELP_LAST},{idx}))'
        ws[f"M{row}"] = f'=IF({idx}=0,"",INDEX(ProjNumbers,{idx}))'
        ws[f"N{row}"] = f'=IF({idx}=0,"",INDEX(ProjCodes,{idx}))'
        ws[f"O{row}"] = f'=IF({idx}=0,"",INDEX($X${HELP_FIRST}:$X${HELP_LAST},{idx}))'
        ws[f"P{row}"] = f'=IF({idx}=0,"",INDEX($Y${HELP_FIRST}:$Y${HELP_LAST},{idx}))'
        ws[f"Q{row}"] = f'=IF({idx}=0,"",INDEX($Z${HELP_FIRST}:$Z${HELP_LAST},{idx}))'
        ws[f"R{row}"] = f'=IF(OR({idx}=0,$L$5=0),"",Q{row}/$L$5)'
    top_last = TOP_FIRST + TOP_N - 1
    style_range(ws, f"L{TOP_FIRST}:R{top_last}", font=f_base, fmt=HOURS)
    style_range(ws, f"R{TOP_FIRST}:R{top_last}", fmt=PCT)
    un, oh, tot = top_last + 1, top_last + 2, top_last + 3
    ws[f"L{un}"] = "Not in project list"
    ws[f"O{un}"] = f"=SUM({DAY('P')})-SUM($X${HELP_FIRST}:$X${HELP_LAST})"
    ws[f"P{un}"] = f"=SUM({DAY('Q')})-SUM($Y${HELP_FIRST}:$Y${HELP_LAST})"
    ws[f"Q{un}"] = f"=O{un}+P{un}"
    ws[f"R{un}"] = f'=IF($L$5=0,"",Q{un}/$L$5)'
    ws[f"L{oh}"], ws[f"M{oh}"] = "Overhead", "Department Hours"
    ws[f"O{oh}"] = f'=IF(OHSystem="NTIME",O5,"")'
    ws[f"P{oh}"] = f'=IF(OHSystem="P6",O5,"")'
    ws[f"Q{oh}"] = "=O5"
    ws[f"R{oh}"] = f'=IF($L$5=0,"",Q{oh}/$L$5)'
    ws[f"L{tot}"] = "Total"
    ws[f"O{tot}"], ws[f"P{tot}"], ws[f"Q{tot}"] = "=M5", "=N5", "=L5"
    style_range(ws, f"L{un}:R{oh}", font=f_base, fmt=HOURS)
    style_range(ws, f"R{un}:R{oh}", fmt=PCT)
    style_range(ws, f"L{tot}:R{tot}", font=f_bold, fill=fill_total, fmt=HOURS)
    ws.conditional_formatting.add(f"L{un}:R{un}", FormulaRule(
        formula=[f"$Q${un}>0"], font=Font(name=FONT, color="C00000", bold=True)))

    # Daily summary ----------------------------------------------------------------------
    section(ws, f"L{DAILY_HEADER - 1}", "DAILY SUMMARY")
    header_row(ws, DAILY_HEADER, 12, ["Date", "Day", "Day Type", "Available", "NTIME", "P6", "Overhead", "Overtime",
                                      "Total", "Status"],
               {"NTIME": fill_ntime, "P6": fill_p6, "Overhead": fill_oh})
    for d in range(1, days + 1):
        row = DAILY_FIRST + d - 1
        date = dt.date(YEAR, m, d)
        ws[f"L{row}"] = date
        ws[f"M{row}"] = f'=TEXT(L{row},"ddd")'
        ws[f"N{row}"] = "Weekend" if date.weekday() >= 5 else "Working"
        ws[f"O{row}"] = "=AvailHours"
        ws[f"P{row}"] = f'=SUMIFS({LOG("I")},{LOG("B")},L{row},{LOG("H")},"NTIME")'
        ws[f"Q{row}"] = f'=SUMIFS({LOG("I")},{LOG("B")},L{row},{LOG("H")},"P6")'
        ws[f"R{row}"] = (f'=IF(N{row}="Overhead only",O{row},'
                         f'IF(AND(N{row}="Working",P{row}+Q{row}>0),MAX(0,O{row}-P{row}-Q{row}),0))')
        ws[f"S{row}"] = f'=IF(OR(N{row}="Working",N{row}="Overhead only"),MAX(0,P{row}+Q{row}-O{row}),0)'
        ws[f"T{row}"] = f"=P{row}+Q{row}+R{row}"
        ws[f"U{row}"] = (
            f'=IF(N{row}="Working",IF(P{row}+Q{row}=0,IF(L{row}<=TODAY(),"Not logged",""),'
            f'IF(P{row}+Q{row}>O{row},IF(AllowOT="Yes","Overtime","Over "&O{row}&" h!"),"✓")),'
            f'IF(N{row}="Overhead only","✓ Overhead day",'
            f'IF(P{row}+Q{row}>0,"Entries on "&LOWER(N{row})&" day",N{row})))'
        )
    style_range(ws, f"L{DAILY_FIRST}:U{d_last}", font=f_base, fmt=HOURS)
    style_range(ws, f"L{DAILY_FIRST}:L{d_last}", fmt="dd-mmm")
    style_range(ws, f"M{DAILY_FIRST}:M{d_last}", align=center)
    for col in "NO":
        style_range(ws, f"{col}{DAILY_FIRST}:{col}{d_last}", font=f_input, fill=fill_input, align=center)
    ws[f"O{DAILY_FIRST}"].comment = Comment(
        "Available hours come from Settings. Type a number here to override a single day.", "Timesheet")
    dv_type = DataValidation(type="list", formula1='"Working,Overhead only,Leave,Public Holiday,Weekend"',
                             showErrorMessage=True)
    dv_avail = DataValidation(type="decimal", operator="between", formula1="0", formula2="24", showErrorMessage=True)
    ws.add_data_validation(dv_type)
    ws.add_data_validation(dv_avail)
    dv_type.add(f"N{DAILY_FIRST}:N{d_last}")
    dv_avail.add(f"O{DAILY_FIRST}:O{d_last}")
    dr = f"L{DAILY_FIRST}:U{d_last}"
    ws.conditional_formatting.add(dr, FormulaRule(
        formula=[f"$L{DAILY_FIRST}=TODAY()"], fill=PatternFill("solid", fgColor="D6E3FF")))
    ws.conditional_formatting.add(dr, FormulaRule(
        formula=[f'OR($N{DAILY_FIRST}="Weekend",$N{DAILY_FIRST}="Leave",$N{DAILY_FIRST}="Public Holiday")'],
        fill=PatternFill("solid", fgColor="F0F0F0"), font=Font(name=FONT, color="808080")))
    ws.conditional_formatting.add(f"U{DAILY_FIRST}:U{d_last}", FormulaRule(
        formula=[f'OR(U{DAILY_FIRST}="Not logged",LEFT(U{DAILY_FIRST},5)="Over ",LEFT(U{DAILY_FIRST},7)="Entries")'],
        font=Font(name=FONT, color="C00000", bold=True), fill=PatternFill("solid", fgColor="FDE2E1")))
    ws.conditional_formatting.add(f"U{DAILY_FIRST}:U{d_last}", FormulaRule(
        formula=[f'LEFT(U{DAILY_FIRST},1)="✓"'], font=Font(name=FONT, color="1A7F37", bold=True)))
    dtot = d_last + 1
    ws[f"L{dtot}"] = "Month total"
    for col in "PQRST":
        ws[f"{col}{dtot}"] = f"=SUM({col}{DAILY_FIRST}:{col}{d_last})"
    style_range(ws, f"L{dtot}:U{dtot}", font=f_bold, fill=fill_total, fmt=HOURS)

    for col in ("W", "X", "Y", "Z", "AA", "AB", "AC", "AD", "AE", "AF", "AG", "AH"):
        ws.column_dimensions[col].hidden = True
    ws.freeze_panes = "A7"
    ws.page_setup.orientation = "landscape"
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True


for month in range(1, 13):
    build_month(month)

# ---- Year dashboard -----------------------------------------------------------------------
yd.sheet_view.showGridLines = False
yd.sheet_properties.tabColor = ORANGE
widths(yd, {"A": 2, "B": 26, "C": 14, "D": 14, "E": 14, "F": 12, "G": 12, "H": 12, "I": 12, "J": 13})
yd["B1"] = f"{YEAR} Dashboard"
yd["B1"].font = f_title
yd["B2"] = "Year totals from the month sheets. Nothing to enter here."
yd["B2"].font = f_sub
section(yd, "B3", "YEAR SUMMARY")
section(yd, "B7", "MONTHLY SUMMARY")
header_row(yd, 8, 2, ["Month", "NTIME", "P6", "Overhead", "Overtime", "Total", "Utilisation", "Days logged",
                      "Days not logged"], {"NTIME": fill_ntime, "P6": fill_p6, "Overhead": fill_oh})
yd["L8"], yd["M8"], yd["N8"] = "NTIME", "P6", "Overhead"  # chart data: project-only system hours
for i, name in enumerate(MONTHS):
    row = 9 + i
    yd[f"B{row}"] = name
    yd[f"B{row}"].hyperlink = f"#'{name}'!A1"
    for col, src in zip("CDEFGHIJ", "MNOPLQRS"):
        yd[f"{col}{row}"] = f"={name}!{src}5"
    yd[f"L{row}"], yd[f"M{row}"], yd[f"N{row}"] = f"={name}!AF16", f"={name}!AG16", f"={name}!AH16"
style_range(yd, "B9:J20", font=f_base, fmt=HOURS)
style_range(yd, "B9:B20", font=Font(name=FONT, size=10, color=NAVY, underline="single"))
style_range(yd, "H9:H20", fmt=PCT)
style_range(yd, "I9:J20", fmt="0")
yd["B21"] = "Year total"
for col in "CDEFGIJ":
    yd[f"{col}21"] = f"=SUM({col}9:{col}20)"
yd["H21"] = '=IF(G21=0,"-",(SUM(L9:L20)+SUM(M9:M20))/G21)'
style_range(yd, "B21:J21", font=f_bold, fill=fill_total, fmt=HOURS)
yd["H21"].number_format = PCT
for col in "IJ":
    yd[f"{col}21"].number_format = "0"

year_kpis = [("Total hours", "=G21", HOURS), ('=IF(OHSystem="NTIME","NTIME (incl. overhead)","NTIME")', "=C21", HOURS),
             ('=IF(OHSystem="P6","P6 (incl. overhead)","P6")', "=D21", HOURS), ("Overhead", "=E21", HOURS),
             ("Overtime", "=F21", HOURS), ("Utilisation", "=H21", PCT), ("Days logged", "=I21", "0"),
             ("Days not logged", "=J21", "0")]
for i, (label, formula, fmt) in enumerate(year_kpis):
    col = get_column_letter(2 + i)
    lc, vc = yd[f"{col}4"], yd[f"{col}5"]
    lc.value, vc.value = label, formula
    lc.font, vc.font = f_kpi_label, f_kpi
    lc.fill = vc.fill = fill_kpi
    lc.alignment = Alignment(horizontal="center", wrap_text=True)
    vc.alignment = center
    vc.number_format = fmt
    lc.border = vc.border = box
yd["C5"].font = Font(name=FONT, size=16, bold=True, color=BLUE)
yd["D5"].font = Font(name=FONT, size=16, bold=True, color=ORANGE)
yd["E5"].font = Font(name=FONT, size=16, bold=True, color="6E7781")
yd.row_dimensions[4].height = 26
yd.row_dimensions[5].height = 30

ychart = BarChart()
ychart.type, ychart.grouping, ychart.overlap = "col", "stacked", 100
ychart.title = f"Hours per month - {YEAR}"
ychart.y_axis.title = "Hours"
ychart.y_axis.majorGridlines = None
ychart.height, ychart.width = 8, 24
ychart.add_data(Reference(yd, min_col=12, max_col=14, min_row=8, max_row=20), titles_from_data=True)
ychart.set_categories(Reference(yd, min_col=2, min_row=9, max_row=20))
for series, color in zip(ychart.series, (BLUE, ORANGE, GREY)):
    series.graphicalProperties.solidFill = color
    series.graphicalProperties.line.solidFill = color
ychart.legend.position = "b"
ychart.visible_cells_only = False
yd.add_chart(ychart, "B23")

section(yd, "B41", f"TOP PROJECTS {YEAR}")
header_row(yd, 42, 2, ["Project", "Project Number", "Oracle Code", "NTIME", "P6", "Total", "% of year"],
           {"NTIME": fill_ntime, "P6": fill_p6})
pkeys = f"Projects!$I${PROJ_FIRST}:$I${PROJ_LAST}"
for k in range(1, 21):
    row = 42 + k
    yd[f"P{row}"] = f"=IF(LARGE({pkeys},{k})<=0,0,MATCH(LARGE({pkeys},{k}),{pkeys},0))"
    idx = f"$P{row}"
    for col, src in zip("BCDEFG", "BCDEFG"):
        yd[f"{col}{row}"] = f'=IF({idx}=0,"",INDEX(Projects!${src}${PROJ_FIRST}:${src}${PROJ_LAST},{idx}))'
    yd[f"H{row}"] = f'=IF(OR({idx}=0,$G$21=0),"",G{row}/$G$21)'
style_range(yd, "B43:H62", font=f_base, fmt=HOURS)
style_range(yd, "H43:H62", fmt=PCT)
for col in ("L", "M", "N", "P"):
    yd.column_dimensions[col].hidden = True
yd.page_setup.fitToWidth = 1
yd.page_setup.fitToHeight = 0
yd.sheet_properties.pageSetUpPr.fitToPage = True

# Sheet order: Guide, Year dashboard, months, Projects, Settings.
order = ["Guide", f"{YEAR} Dashboard"] + MONTHS + ["Projects", "Settings"]
wb._sheets = [wb[n] for n in order]
for ws in wb.worksheets:
    for row in ws.iter_rows():
        for c in row:
            if c.font is None or c.font.name != FONT:
                c.font = Font(name=FONT, size=c.font.size if c.font else 10, bold=c.font.bold if c.font else False,
                              italic=c.font.italic if c.font else False, color=c.font.color if c.font else None)
current = dt.date.today()
wb.active = order.index(MONTHS[current.month - 1]) if current.year == YEAR else 1
wb.save(OUT)
print(f"Saved {OUT}")
